"""Extract paragraph/table text from a .docx into a UTF-8 text file.

Usage: python -I tools/extract_docx.py <input.docx> <output.txt>
"""
import sys
import docx

src, dst = sys.argv[1], sys.argv[2]
doc = docx.Document(src)
lines = []
for para in doc.paragraphs:
    text = para.text.strip()
    if text:
        style = para.style.name if para.style is not None else ""
        lines.append(f"[{style}] {text}")

for idx, table in enumerate(doc.tables, 1):
    lines.append(f"\n=== TABLE {idx} ===")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        lines.append(" | ".join(cells))

with open(dst, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print(f"paragraphs/tables lines: {len(lines)}, chars: {sum(len(x) for x in lines)}")
